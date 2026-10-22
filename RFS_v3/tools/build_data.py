"""Builds client-side lookup files from the founder's Excel libraries.
Run from the repo root:  MCC_LIB_DIR=/path/to/libraries python3 tools/build_data.py
Re-run whenever the Recruiter Library or B-School Library is updated."""
import json, re
from openpyxl import load_workbook
import os
P=os.environ.get('MCC_LIB_DIR','./').rstrip('/')+'/'   # folder holding the .xlsx libraries and the JD Library context .md
OUT=os.environ.get('MCC_OUT_DIR','career-clinic-resume-checker/data').rstrip('/')+'/'
CODES={'CON','FIN','IB','SAL','MKT','PM','HR','GM','OPS','SCM','PRC','LEG','BA','ITS','CSM','ENT','HCM','AGB','RTL','RE','PPD','ESG','ECM','MED','ENR','INF','TEL','EDT'}
# Short names that are also common English words or very short: match case-sensitively and only on role-header lines.
COMMON_WORDS={'Target','Oracle','Apple','Mars','Titan','Meta','Navi','Lava','Puma','Visa','Vivo','Ola','SAIL','Amazon','Google','Groww','Acko','boAt','IndiGo','Uber','Zoho','Dell','Citi','Aon','HP','STL','MPL','FIS','ICF','TCI','UKG','ZEE','BSE','NSE','LIC','UPL','DBS','UBS','BNY','EXL','WNS','HGS','OYO','CRED'}
EXTRA_ALIASES={'Hindustan Unilever':['HUL','Unilever India'],'Procter & Gamble':['P&G','PnG','Procter and Gamble'],
 'TCS':['Tata Consultancy Services'],'ITC':['ITC Limited','ITC Ltd'],'Asian Paints':[],'Larsen & Toubro':['L&T'],
 'Deloitte India':['Deloitte'],'KPMG India':['KPMG'],'PwC India':['PwC','PricewaterhouseCoopers'],
 'EY-Parthenon / EY India':['EY','Ernst & Young','EY-Parthenon'],'Accenture':['Accenture Strategy'],
 'Reliance Industries':['RIL'],'Mahindra Group':['Mahindra & Mahindra','M&M'],'Nestle India':['Nestlé','Nestle'],
 'Tech Mahindra':['TechM'],'Kotak Mahindra Bank':['Kotak'],'ICICI Bank':['ICICI'],'Bajaj Finance':['Bajaj Finserv'],
 'Goldman Sachs':['GS'],'J.P. Morgan':['JPMorgan','JP Morgan','JPMC'],'Morgan Stanley':[],'McKinsey & Company':['McKinsey'],
 'Bain & Company':['Bain'],'BCG':['Boston Consulting Group'],'Amazon':['Amazon India'],'Flipkart':[]}
def aliases(name, legal):
    out={name}
    for part in re.split(r'\s*/\s*|\s*\(|\)', name):
        p=part.strip()
        if len(p)>=2: out.add(p)
    out.update(EXTRA_ALIASES.get(name,[]))
    if legal:
        core=re.sub(r'\(.*?\)','',legal)
        core=re.sub(r'\b(Private|Pvt\.?|Limited|Ltd\.?|LLP|India|Inc\.?|Services|Solutions|Technologies|Corporation|Company|Co\.)\b','',core,flags=re.I)
        core=re.sub(r'\s+',' ',core).strip(' ,.-')
        if 4<=len(core)<=40 and len(core.split())<=4: out.add(core)
    STOP={'india','group','limited','global','services','solutions','technologies','capital','bank','finance','consulting','foods','motors',
          'industries','the','and','international','holdings','ventures','labs','digital','retail','energy','power','health','healthcare','media','new','first'}
    return sorted(a for a in out if len(a)>=2 and a.lower() not in STOP)

wb=load_workbook(P+'MyCareerCompany_MBA_Recruiter_Library_v2_1.xlsx',read_only=True)
ws=wb['Company Library']; comps=[]
for i,r in enumerate(ws.iter_rows(values_only=True)):
    if i<7 or not r[1]: continue
    codes=[c.strip() for c in str(r[8] or '').split(';') if c.strip() in CODES]
    top3=[c.strip() for c in str(r[10] or '').split(';') if c.strip() in CODES]
    ind=str(r[5] or '').split(' ')[0]; ind=ind if ind in CODES else None
    comps.append({'n':r[1],'a':aliases(r[1],r[2]),'s':r[4],'i':ind,'p':r[6],'c':codes,'t':top3,
                  'cs':1 if (r[1] in COMMON_WORDS or len(r[1])<=4) else 0})
wb=load_workbook(P+'MyCareerCompany_India_BSchool_Library_Oct2026.xlsx',read_only=True)
ws=wb['Master Library']; schools=[]
MAPD={'consulting':['CON'],'finance':['FIN'],'bfsi':['FIN','IB'],'banking':['FIN'],'investment banking':['IB'],'marketing':['MKT'],
 'sales':['SAL'],'hr':['HR'],'labour':['HR'],'organisation development':['HR'],'business analytics':['BA'],'analytics':['BA'],'data science':['BA'],
 'operations':['OPS'],'supply chain':['SCM'],'logistics':['SCM'],'product':['PM'],'tech':['PM','ITS'],'it':['ITS'],'systems':['ITS'],
 'general management':['GM'],'entrepreneurship':['ENT'],'family business':['ENT'],'agribusiness':['AGB'],'rural':['AGB'],'agri':['AGB'],
 'healthcare':['HCM'],'hospital':['HCM'],'pharma':['HCM'],'public health':['HCM'],'retail':['RTL'],'fmcg':['MKT','SAL'],'communications':['PRC'],
 'advertising':['MKT','MED'],'media':['MED'],'brand':['MKT'],'development':['PPD'],'policy':['PPD'],'international business':['GM','SCM'],
 'infrastructure':['INF'],'energy':['ENR'],'telecom':['TEL'],'real estate':['RE'],'e-commerce':['ECM'],'category':['ECM','RTL'],'manufacturing':['OPS']}
def dom_codes(txt):
    out=[]
    for part in (txt or '').lower().replace('&',';').replace('/',';').split(';'):
        p=part.strip()
        for k,v in MAPD.items():
            if re.search(r'\b'+re.escape(k)+r'\b',p):
                for c in v:
                    if c not in out: out.append(c)
    return out[:4]
for i,r in enumerate(ws.iter_rows(values_only=True)):
    if i<7 or not r[2]: continue
    basis=str(r[14] or '')
    evidence= not basis.upper().startswith('ASSUMED')   # 86% of rows carry an ASSUMED default tag; never use those as a domain signal
    tier={'Tier 1':1,'Tier 2':2,'Tier 3':3}.get(r[10],3)
    al={r[1],r[2]}
    m=re.match(r'IIM\s+(\w+)',str(r[1] or ''))
    if m:
        city=m.group(1); al|={'IIM-'+city,'IIM '+city,'Indian Institute of Management '+city,'Indian Institute of Management, '+city}
        if city in ('Ahmedabad','Bangalore','Calcutta','Lucknow','Kozhikode','Indore'): al.add('IIM'+city[0])
    schools.append({'n':r[1],'a':sorted(a for a in al if a and len(a)>=3),'t':tier,'d':dom_codes(r[13]) if evidence else [],'e':1 if evidence else 0})
hdr='/* AUTO-GENERATED by build_data.py from MyCareerCompany Excel libraries (Oct 2026). Do not hand-edit. */\n'
open(OUT+'rfs-v3-companies.js','w').write(hdr+"(function(r){var D=%s;if(typeof module==='object'&&module.exports){module.exports=D;}else{r.RFS_COMPANIES=D;}})(this);\n"%json.dumps(comps,ensure_ascii=False,separators=(',',':')))
open(OUT+'rfs-v3-schools.js','w').write(hdr+"(function(r){var D=%s;if(typeof module==='object'&&module.exports){module.exports=D;}else{r.RFS_SCHOOLS=D;}})(this);\n"%json.dumps(schools,ensure_ascii=False,separators=(',',':')))
# The JD library itself (jd_library_v2.json already ships in the repo). Copy for local tests only.
lines=open(P+'MyCareerCompany_JD_Library_v2_Antigravity_Context.md',encoding='utf-8').read().split('\n')
recs=[json.loads(l.strip().rstrip(',')) for l in lines if l.strip().startswith('{"id"')]
os.makedirs('test',exist_ok=True); json.dump(recs,open('test/jd_library_v2.json','w'),ensure_ascii=False)
print(len(comps),'companies;',len(schools),'schools (',sum(s['e'] for s in schools),'with evidence-based domain tags );',len(recs),'JDs')
